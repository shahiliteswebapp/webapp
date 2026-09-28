"""
Merge the client's decorative source sheets into src/lib/catalog-data.json.

  ../rawdata/Geo Liting Part 2 - Product Details.xlsx  wall + mirror lamps, with photo file names
  ../rawdata/chandeliers chart.xlsx                     client tagging: product type, style, mounting

Run after scripts/parse-geo-catalogs.py. Idempotent: re-running rebuilds the
tags and Part 2 rows from the sheets.
Photos themselves are converted/uploaded by scripts/upload-catalog-images.mjs;
this script only records their (webp) file names on each system.

Run from the repo root:  python scripts/merge-decorative.py
"""

import json
import os
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT.parent / "rawdata"
CATALOG = ROOT / "src" / "lib" / "catalog-data.json"

PART2_SOURCE = "Geo Liting Wall Lamp & Mirror Lamps Part 2"


def norm_code(code) -> str:
    return re.sub(r"[\s\-]+", "", str(code or "")).upper()


def chart_codes(raw):
    """A chart cell can name several variants: "5218 A,B", "1025-400,600,800",
    "2394-1 SM,AMB", "20188/1". Yield every spelling worth trying."""
    raw = str(raw or "").strip()
    yield raw
    for part in re.split(r"/", raw):
        yield part
    m = re.match(r"^(.*?[\-\s])([^\-\s,]+(?:\s*,\s*[^\-\s,]+)+)$", raw)
    if m:
        stem = m.group(1)
        for v in m.group(2).split(","):
            yield stem + v.strip()


def alias_keys(code):
    """Lookup keys for a catalogue code: as printed, without a parenthetical
    note ("M406-AB (A)", "FY03-AB (ADJUSTABLE)"), and before a variant list
    ("2212-3,A,B,C")."""
    code = str(code or "")
    bare = re.sub(r"\(.*?\)", "", code)
    return {norm_code(code), norm_code(bare), norm_code(bare.split(",")[0])}


def find_items(by_code, decorative, raw):
    """Catalogue items a chart row refers to: an exact / alias match, else the
    whole family when the chart names a stem ("B6001" -> B6001-S, B6001-M)."""
    for c in chart_codes(raw):
        hit = by_code.get(norm_code(c))
        if hit:
            return [hit]
    stem = str(raw or "").strip()
    if len(norm_code(stem)) < 4:
        return []
    fam = [
        c for c in decorative
        if re.match(re.escape(stem) + r"[\s\-(]", str(c["sourceCode"]), re.I)
    ]
    return fam


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


# ---- chart vocab -> app vocab ------------------------------------------------

TYPE_MAP = {
    "PENDENT": "Hanging light",
    "PENDANT": "Hanging light",
    "CHANDELIER": "Chandelier",
    "PENDENT/CHANDELIER": "Chandelier",
    "CHANDELIER/PENDANT": "Chandelier",
    "WALL LIGHT": "Wall light",
    "WALL LITGH": "Wall light",
    "WALL LIGHT/MIRROR LIGHT": "Mirror light",
    "BED SIDE LIGHT": "Bed side light",
    "PICTURE LIGHT": "Picture light",
    "TABLE LAMP": "Table / floor lamp",
    "FLOOR LAMP": "Table / floor lamp",
    "FLOOR TABLE LAMP": "Table / floor lamp",
    "CEILING LIGHT": "Ceiling light",
}

STYLE_MAP = {
    "MORDERN": ["modern"],
    "MORDEN": ["modern"],
    "MODERN": ["modern"],
    "CLASICAL": ["classical"],
    "CLASSICAL": ["classical"],
    "MORDERN/CLASICAL": ["modern", "classical"],
    "CRYSTAL": ["crystal"],
}

MOUNT_WORDS = {
    "HANGING": "Hanging",
    "WALL MOUNTED": "Wall mounted",
    "CEILING MOUNT": "Ceiling mounted",
    "FLOOR MOUNTED": "Floor mounted",
    "LINEAR": "Linear",
    "LENEAR": "Linear",
    "CORNER": "Corner",
    "CONER": "Corner",
    "SINGLE": "Single",
    "SIN": "Single",
    "SL": "Single",
    "S": "Single",
    "STAIRCASE": "Staircase",
    "STAIR CASE": "Staircase",
    "DOUBLE HEIGHT": "Double height",
    "OUTDOOR": "Outdoor",
}


def map_type(raw):
    key = re.sub(r"\s+", " ", str(raw or "")).strip().upper()
    return TYPE_MAP.get(key)


def map_styles(raw):
    key = re.sub(r"\s+", "", str(raw or "")).upper()
    return STYLE_MAP.get(key, [])


def map_mountings(raw):
    out = []
    for part in str(raw or "").split("/"):
        key = re.sub(r"\s+", " ", part).strip().upper()
        tag = MOUNT_WORDS.get(key)
        if tag and tag not in out:
            out.append(tag)
    return out


def webp(name: str) -> str:
    return os.path.splitext(name.strip())[0] + ".webp"


def main():
    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))

    # Drop previous Part 2 rows so re-runs don't duplicate them.
    catalog = [c for c in catalog if c.get("source") != PART2_SOURCE]

    by_code = {}
    for c in catalog:
        if c["kind"] == "decorative":
            for k in alias_keys(c["sourceCode"]):
                by_code.setdefault(k, c)

    # ---- Part 2 (wall / mirror lamps, with photos) ----
    wb = openpyxl.load_workbook(RAW / "Geo Liting Part 2 - Product Details.xlsx", data_only=True)
    added = 0
    for r in wb["Products"].iter_rows(min_row=2, values_only=True):
        _, _page, code, lamp, size, finish, material, sku, price, own, shared, _notes = r[:12]
        if not code:
            continue
        images = [webp(x) for x in f"{own or ''},{shared or ''}".split(",") if x.strip()]
        existing = by_code.get(norm_code(code))
        if existing:
            have = existing.setdefault("images", [])
            have.extend(i for i in images if i not in have)
            continue
        code = str(code).strip()
        sysid = "geo-" + re.sub(r"[^a-z0-9]+", "-", code.lower()).strip("-")
        led = str(lamp or "").strip().lower() == "led"
        item = {
            "id": sysid,
            "sourceCode": code,
            "name": f"Wall light {code}",
            "kind": "decorative",
            "decorType": "Wall light",
            "mounting": None,
            "style": None,
            "indoorOutdoor": None,
            "unit": "nos",
            "unitCost": float(price) if price else 0,
            "size": clean(size),
            "finish": clean(finish),
            "material": clean(material),
            "lamp": clean(lamp),
            "sku": clean(sku),
            "automatic": led,
            "interfaceOptions": (
                [
                    {"interface": "DIMMABLE", "control": "dimmable", "price": None},
                    {"interface": "TUNABLE", "control": "tunable", "price": None},
                ]
                if led
                else []
            ),
            "rules": [],
            "source": PART2_SOURCE,
            "images": images,
        }
        catalog.append(item)
        by_code[norm_code(code)] = item
        added += 1

    # ---- client tagging chart ----
    wb = openpyxl.load_workbook(RAW / "chandeliers chart.xlsx", data_only=True)
    tagged = 0
    unmatched = []
    decorative = [c for c in catalog if c["kind"] == "decorative"]
    for r in wb.active.iter_rows(min_row=2, values_only=True):
        code, ptype, style, mount = r[1], r[3], r[4], r[5]
        hits = find_items(by_code, decorative, code)
        if not hits:
            unmatched.append(str(code))
            continue
        t = map_type(ptype)
        styles = map_styles(style)
        mounts = map_mountings(mount)
        for c in hits:
            if t:
                c["decorType"] = t
                c["name"] = f"{t} {c['sourceCode']}"
            c["styleTags"] = styles
            c["style"] = " / ".join(styles) or None
            c["mountingTags"] = [m for m in mounts if m != "Outdoor"]
            c["mounting"] = " / ".join(c["mountingTags"]) or None
            c["indoorOutdoor"] = "outdoor" if "Outdoor" in mounts else "indoor"
            tagged += 1

    for c in catalog:
        if c["kind"] == "decorative":
            c.setdefault("styleTags", [])
            c.setdefault("mountingTags", [])
            c.setdefault("images", [])

    CATALOG.write_text(json.dumps(catalog, ensure_ascii=False), encoding="utf-8")
    dec = [c for c in catalog if c["kind"] == "decorative"]
    print(f"part2 added {added}, tagged {tagged}, decorative {len(dec)}, "
          f"with images {sum(1 for c in dec if c['images'])}, total {len(catalog)}")
    print(f"chart codes with no catalogue item ({len(unmatched)}): {unmatched}")


if __name__ == "__main__":
    main()
